"""P3 tests for multi-format extractors & mount scanning (CONTRACTS §C2).

Uses real binary fixtures generated under tests/fixtures/docs (mirroring the
doxtr/dox reference project's real-document approach).
"""

from __future__ import annotations

import os
from pathlib import Path

import pytest

from doxtr_rag.extractors import (
    extract_file,
    scan_mount,
    supported_extensions,
)
from doxtr_rag.schema import AccessControl, DocumentChunk

FIXTURES = Path(__file__).parent / "fixtures" / "docs"
SHARED = AccessControl.SHARED_ORG


def _chunks(name: str, scope: AccessControl = SHARED) -> list[DocumentChunk]:
    return list(extract_file(str(FIXTURES / name), scope))


def test_all_expected_extensions_registered() -> None:
    assert supported_extensions() >= {
        ".pdf",
        ".xlsx",
        ".ods",
        ".docx",
        ".odt",
        ".pptx",
        ".keynote",
    }


# -- PDF ---------------------------------------------------------------------
def test_pdf_page_level_provenance() -> None:
    chunks = _chunks("sample.pdf")
    assert len(chunks) >= 1
    assert chunks[0].metadata["provenance"] == {"page": 1}
    assert "Hello RAG PDF" in chunks[0].text


def test_corrupt_pdf_skipped_gracefully() -> None:
    # Must not raise; yields nothing.
    assert _chunks("corrupt.pdf") == []


# -- Spreadsheets ------------------------------------------------------------
def test_xlsx_markdown_table_and_cell_provenance() -> None:
    chunks = _chunks("sample.xlsx")
    data = next(c for c in chunks if "Widget" in c.text)
    assert "| Name | Qty |" in data.text
    assert "cell" in data.metadata["provenance"]


def test_ods_markdown_table() -> None:
    chunks = _chunks("sample.ods")
    assert any("| A | B |" in c.text for c in chunks)
    assert all("cell" in c.metadata["provenance"] for c in chunks)


# -- Presentations -----------------------------------------------------------
def test_pptx_combines_title_body_notes() -> None:
    chunks = _chunks("sample.pptx")
    assert len(chunks) == 1
    text = chunks[0].text
    assert "Slide Title" in text
    assert "Bullet body text" in text
    assert "speaker notes" in text.lower()
    assert chunks[0].metadata["provenance"] == {"slide": 1}


def test_keynote_best_effort_skips(caplog: pytest.LogCaptureFixture) -> None:
    with caplog.at_level("WARNING"):
        chunks = _chunks("sample.keynote")
    assert chunks == []
    assert any("keynote" in r.message.lower() for r in caplog.records)


# -- Word / OpenOffice -------------------------------------------------------
def test_docx_heading_aware() -> None:
    chunks = _chunks("sample.docx")
    titles = [c.metadata["section_title"] for c in chunks]
    assert "Section One" in titles
    # Breadcrumb includes the file basename + heading trail.
    sec = next(c for c in chunks if c.metadata["section_title"] == "Section One")
    assert "sample.docx" in sec.metadata["breadcrumb"]


def test_odt_heading_aware() -> None:
    chunks = _chunks("sample.odt")
    joined = " ".join(c.text for c in chunks)
    assert "Body under title." in joined
    assert "Sub body." in joined


def test_docx_xml_fallback_recovers_non_paragraph_body(tmp_path: Path) -> None:
    """Body text python-docx does not surface as paragraphs is still recovered.

    Builds a .docx whose only ``doc.paragraphs`` entry is a section label
    (custom/Normal style, so no real Heading styles exist) while the actual
    body text lives inside a table cell -- which ``doc.paragraphs`` does NOT
    surface. This forces ``extract_docx`` down the raw-XML ``<w:t>`` fallback
    (``_extract_docx_xml``). The test asserts the body text is recovered from
    the raw XML and that the section label is preserved as heading structure.

    Regression guard: if the XML fallback is removed and only the primary
    ``doc.paragraphs`` pass runs, the table-cell body text is lost entirely
    and this test fails.
    """
    from docx import Document

    doc = Document()
    # Section label paragraph (Normal style -> not a recognized Heading).
    doc.add_paragraph("Section One")
    # Body text inside a table cell: NOT returned by doc.paragraphs, but it is
    # present as a <w:t> node in word/document.xml.
    table = doc.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "This body text lives in a table cell."
    docx_path = tmp_path / "table_body.docx"
    doc.save(str(docx_path))

    chunks = list(extract_file(str(docx_path), SHARED))

    # Body text recovered via the raw-XML fallback.
    joined = " ".join(c.text for c in chunks)
    assert "This body text lives in a table cell." in joined
    # Section structure preserved: the label becomes a heading with breadcrumb.
    titles = [c.metadata["section_title"] for c in chunks]
    assert "Section One" in titles
    sec = next(c for c in chunks if c.metadata["section_title"] == "Section One")
    assert "table_body.docx" in sec.metadata["breadcrumb"]
    assert "Section One" in sec.metadata["breadcrumb"]


def test_docx_xml_fallback_preserves_real_heading_with_table_body(
    tmp_path: Path,
) -> None:
    """A real Heading style is preserved (not re-promoted) on the XML fallback.

    Builds a .docx with a REAL ``Heading 1`` styled paragraph plus body text
    that lives ONLY inside a table cell (``doc.paragraphs`` surfaces the
    heading but no body paragraph). ``has_body`` is therefore False and
    ``extract_docx`` takes the raw-XML fallback with ``any_real_heading=True``.
    Asserts (a) the heading keeps its API-derived level (it is recognised as a
    heading and produces breadcrumb structure) and (b) the table-cell body is
    recovered under that heading.

    Regression guard for the branch where real headings coexist with body that
    only python-docx's paragraph API misses.
    """
    from docx import Document

    doc = Document()
    doc.add_heading("Real Heading", level=1)
    table = doc.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "Body only in a table cell."
    docx_path = tmp_path / "heading_table.docx"
    doc.save(str(docx_path))

    chunks = list(extract_file(str(docx_path), SHARED))

    joined = " ".join(c.text for c in chunks)
    assert "Body only in a table cell." in joined
    titles = [c.metadata["section_title"] for c in chunks]
    assert "Real Heading" in titles
    sec = next(c for c in chunks if c.metadata["section_title"] == "Real Heading")
    # Heading preserved as structure: it appears in the breadcrumb trail and
    # the table-cell body is grouped under it.
    assert "Real Heading" in sec.metadata["breadcrumb"]
    assert "Body only in a table cell." in sec.text


def test_docx_xml_fallback_recognizes_split_run_heading(tmp_path: Path) -> None:
    """A heading label split across multiple runs is still recognised.

    python-docx (and Word) split a paragraph across several ``<w:r>``/``<w:t>``
    runs at formatting boundaries. If the fallback matched bare ``<w:t>`` nodes
    it would never match a heading label stored as several runs, silently
    degrading the heading to body text. This builds a Heading 1 paragraph whose
    text is written as two separately-formatted runs ("Split " + "Heading",
    the second bold) plus table-cell body, forcing the XML fallback, and
    asserts the reconstructed paragraph is recognised as a heading.
    """
    from docx import Document

    doc = Document()
    heading = doc.add_heading(level=1)
    heading.add_run("Split ")
    heading.add_run("Heading").bold = True  # forces a separate run/<w:t>
    table = doc.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "Body under a split heading."
    docx_path = tmp_path / "split_heading.docx"
    doc.save(str(docx_path))

    chunks = list(extract_file(str(docx_path), SHARED))

    titles = [c.metadata["section_title"] for c in chunks]
    # The two runs reconstruct to "Split Heading" and are recognised as a
    # heading (NOT split into / degraded to body text).
    assert "Split Heading" in titles
    sec = next(c for c in chunks if c.metadata["section_title"] == "Split Heading")
    assert "Split Heading" in sec.metadata["breadcrumb"]
    assert "Body under a split heading." in sec.text


# -- Private / password routing ----------------------------------------------
def test_private_scope_preserved() -> None:
    chunks = _chunks("sample.pdf", AccessControl.LOCAL_PRIVATE)
    assert all(c.access_scope is AccessControl.LOCAL_PRIVATE for c in chunks)


# -- Mount scanning ----------------------------------------------------------
def test_scan_mount_finds_all(tmp_path: Path) -> None:
    # Copy a couple of fixtures into a scratch mount.
    import shutil

    mount = tmp_path / "mnt"
    mount.mkdir()
    for name in ("sample.pdf", "sample.xlsx"):
        shutil.copy(FIXTURES / name, mount / name)
    chunks = list(scan_mount(str(mount), SHARED))
    sources = {os.path.basename(c.metadata["source_file"]) for c in chunks}
    assert {"sample.pdf", "sample.xlsx"} <= sources


def test_scan_missing_mount_warns(caplog: pytest.LogCaptureFixture) -> None:
    with caplog.at_level("WARNING"):
        chunks = list(scan_mount("/nonexistent/mount/xyz", SHARED))
    assert chunks == []
    assert any("not found" in r.message for r in caplog.records)


def test_scan_skips_doxtr_dir(tmp_path: Path) -> None:
    import shutil

    mount = tmp_path / "mnt"
    (mount / ".doxtr").mkdir(parents=True)
    shutil.copy(FIXTURES / "sample.pdf", mount / ".doxtr" / "sample.pdf")
    chunks = list(scan_mount(str(mount), SHARED))
    assert chunks == []


def test_unknown_extension_ignored(tmp_path: Path) -> None:
    f = tmp_path / "note.txt"
    f.write_text("hi")
    assert list(extract_file(str(f), SHARED)) == []


def test_scan_mount_blocks_symlink_escape(
    tmp_path: Path, caplog: pytest.LogCaptureFixture
) -> None:
    import shutil

    mount = tmp_path / "shared_mount"
    mount.mkdir()
    outside = tmp_path / "private"
    outside.mkdir()
    secret = outside / "secret.pdf"
    shutil.copy(FIXTURES / "sample.pdf", secret)
    link = mount / "leak.pdf"
    try:
        os.symlink(secret, link)
    except (OSError, NotImplementedError):
        pytest.skip("symlinks unsupported")
    with caplog.at_level("WARNING"):
        chunks = list(scan_mount(str(mount), SHARED))
    # The symlink escaping the mount is skipped fail-closed -> no shared chunk.
    assert chunks == []
    assert any("escapes mount" in r.message for r in caplog.records)


def test_keynote_and_key_both_registered() -> None:
    assert {".keynote", ".key"} <= supported_extensions()


# -- error paths / password protection ---------------------------------------
def test_encrypted_pdf_routes_private_and_skips(
    caplog: pytest.LogCaptureFixture,
) -> None:
    with caplog.at_level("INFO"):
        chunks = _chunks("encrypted.pdf")
    # Cannot read protected content -> skipped (never leaked to shared store).
    assert chunks == []
    assert any("LOCAL_PRIVATE" in r.message for r in caplog.records)


@pytest.mark.parametrize(
    "name", ["bad.xlsx", "bad.docx", "bad.pptx", "bad.ods", "bad.odt"]
)
def test_corrupt_office_files_skip_gracefully(name: str) -> None:
    # Must not raise; yields nothing.
    assert _chunks(name) == []


def test_empty_spreadsheet_sheet_skipped() -> None:
    # sample.xlsx has an "Empty" sheet with no rows -> no chunk for it.
    chunks = _chunks("sample.xlsx")
    assert all(c.metadata["section_title"] != "Empty" for c in chunks)


def test_sha256_file_helper() -> None:
    from doxtr_rag.extractors import _sha256_file

    h = _sha256_file(str(FIXTURES / "sample.pdf"))
    assert len(h) == 64 and all(c in "0123456789abcdef" for c in h)


def test_matrix_to_markdown_empty() -> None:
    from doxtr_rag.extractors import _matrix_to_markdown

    assert _matrix_to_markdown([]) == ""
